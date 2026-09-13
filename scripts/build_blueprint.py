"""Build the editorial outline and editable review document from maintained sources.

Requires python-docx. This builds documentation only; it does not execute book labs.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DATA = json.loads((ROOT / "chapters.json").read_text())
SOURCES = json.loads((ROOT / "research/sources.json").read_text())
SOURCE_MAP = {s["id"]: s for s in SOURCES}


def validate():
    chapters = DATA["chapters"]
    assert [c["number"] for c in chapters] == list(range(1, 31))
    assert [p["number"] for p in DATA["parts"]] == list(range(1, 9))
    assert len(SOURCE_MAP) == len(SOURCES)
    assert sum(c["pages"] for c in chapters) == 500
    for part in DATA["parts"]:
        selected = [c for c in chapters if c["part"] == part["number"]]
        assert selected and sum(c["pages"] for c in selected) == part["pages"]
    for c in chapters:
        assert c["pages"] > 0 and len(c["topics"]) == 4
        assert all(0 < n < c["number"] for n in c["prerequisites"])
        assert c["sources"] and all(s in SOURCE_MAP for s in c["sources"])
        for key in ("title", "transfer", "workshop", "failure", "completion"):
            assert c[key].strip(), (c["number"], key)
    for source in SOURCES:
        assert source["url"].startswith("https://")
        for key in ("title", "author", "date", "claim", "limitation", "access_note"):
            assert source[key].strip()


def outline_text():
    lines = ["# Detailed chapter outline", "",
             "This outline specifies 30 chapters in eight parts. Chapter page allocations are estimates for the future manuscript and total 500 pages. They exclude front matter, references, index, and appendices.", "",
             "Each brief preserves the approved chapter topics and adds the software-engineering connection, prerequisites, a workshop, a failure exercise, and assessable completion evidence. Reading links are sources for the concepts; the proposed exercises are original editorial designs.", "",
             "| Part | Chapters | Main-text pages |", "| --- | --- | --- |"]
    for p in DATA["parts"]:
        cs = [c for c in DATA["chapters"] if c["part"] == p["number"]]
        lines.append(f'| {p["number"]}. {p["title"]} | {cs[0]["number"]}–{cs[-1]["number"]} | {p["pages"]} |')
    lines += ["| Total | 30 chapters | 500 |", ""]
    for p in DATA["parts"]:
        lines += [f'## Part {p["number"]} — {p["title"]}', "", f'Estimated main text: {p["pages"]} pages.', ""]
        for c in [c for c in DATA["chapters"] if c["part"] == p["number"]]:
            prereqs = ", ".join(str(n) for n in c["prerequisites"]) or "Existing software-engineering experience"
            lines += [f'### Chapter {c["number"]} — {c["title"]}', "",
                      f'**Page allocation:** {c["pages"]}. **Prerequisites:** {prereqs}.', "",
                      f'**Software-engineering connection:** {c["transfer"]}', "", "**Topics and learning outcomes**", ""]
            lines += [f"- {topic}" for topic in c["topics"]]
            lines += ["", f'**Workshop:** {c["workshop"]}', "", f'**Failure exercise:** {c["failure"]}', "",
                      f'**Completion evidence:** {c["completion"]}', ""]
            refs = [f'[{SOURCE_MAP[s]["title"]}]({SOURCE_MAP[s]["url"]})' for s in c["sources"]]
            lines += ["**Further reading:** " + "; ".join(refs) + ".", ""]
    return "\n".join(lines)


def source_notes():
    lines = ["# Source notes", "",
             "The notes below record publication context and the limits of the source's contribution. All sources were consulted on September 5, 2026. They support the curriculum's concepts; they do not validate the proposed page counts, exercises, or acceptance targets.", ""]
    for s in SOURCES:
        lines += [f'## {s["title"]}', "", f'{s["author"]}. {s["date"]}. {s["kind"]}.', "",
                  f'[{s["title"]}]({s["url"]})', "", f'**Contribution:** {s["claim"]}', "",
                  f'**Boundary:** {s["limitation"]}', ""]
    return "\n".join(lines)


def append_inline(p, text):
    # The source documents use only inline links and bold emphasis.
    pattern = re.compile(r"(\[[^\]]+\]\(https://[^)]+\)|\*\*[^*]+\*\*)")
    for bit in pattern.split(text):
        if not bit:
            continue
        if bit.startswith("[") and "](" in bit:
            label, url = bit[1:-1].split("](", 1)
            link = OxmlElement("w:hyperlink")
            rel = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
            link.set(qn("r:id"), rel)
            run = OxmlElement("w:r")
            props = OxmlElement("w:rPr")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "245A81")
            props.append(color)
            run.append(props)
            txt = OxmlElement("w:t")
            txt.text = label
            run.append(txt)
            link.append(run)
            p._p.append(link)
        else:
            run = p.add_run(bit[2:-2] if bit.startswith("**") else bit)
            run.bold = bit.startswith("**")


def add_table(doc, lines):
    rows = [[x.strip() for x in line.strip().strip("|").split("|")] for line in lines]
    rows = [r for r in rows if not all(re.fullmatch(r":?-+:?", x) for x in r)]
    ncol = len(rows[0])
    table = doc.add_table(rows=0, cols=ncol)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [2.9, 1.0, 3.1] if ncol == 3 else [7.0 / ncol] * ncol
    if "Part" in rows[0][0]:
        widths = [4.5, 1.0, 1.5]
    elif rows[0][0] == "Behavior":
        widths = [1.4, 2.4, 3.2]
    elif rows[0][0] == "Scenario":
        widths = [2.35, 0.85, 3.8]
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            element.set(qn(f"w:{key}"), value)
        borders.append(element)
    table._tbl.tblPr.append(borders)
    for i, values in enumerate(rows):
        row = table.add_row()
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        if i == 0:
            header = OxmlElement("w:tblHeader")
            row._tr.get_or_add_trPr().append(header)
        for j, value in enumerate(values):
            cell = row.cells[j]
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                el = OxmlElement(f"w:{side}")
                el.set(qn("w:w"), "90")
                el.set(qn("w:type"), "dxa")
                margins.append(el)
            tcpr.append(margins)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "E7EDF3" if i == 0 else "FFFFFF")
            tcpr.append(shade)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.05
            if j == 1 and ("Chapters" in rows[0][j] or "Part" in rows[0][0]):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            append_inline(p, value)
            for run in p.runs:
                run.font.size = Pt(9)
                run.bold = i == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_markdown(doc, text):
    lines = text.splitlines()
    keep_records = text.startswith("# Source notes")
    record_start = None

    def finish_record():
        if record_start is not None:
            record = doc.paragraphs[record_start:]
            for paragraph in record[:-1]:
                paragraph.paragraph_format.keep_with_next = True
            if record:
                record[-1].paragraph_format.keep_with_next = False

    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            add_table(doc, rows)
            continue
        match = re.match(r"^(#{1,3}) (.+)$", line)
        if match:
            finish_record()
            record_start = None
            if match[2].startswith("Chapter ") or (keep_records and len(match[1]) == 2):
                record_start = len(doc.paragraphs)
            p = doc.add_heading(match[2], level=len(match[1]))
            if len(match[1]) == 1:
                p.paragraph_format.page_break_before = True
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            append_inline(p, line[2:])
        else:
            chunk = [line]
            while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(r"^(#|\||- )", lines[i + 1]):
                i += 1
                chunk.append(lines[i].strip())
            p = doc.add_paragraph()
            append_inline(p, " ".join(chunk))
            if line.startswith("**Topics") or line.startswith("Estimated main text:"):
                p.paragraph_format.keep_with_next = True
        i += 1
    finish_record()


def create_doc(outline):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.75)
    section.footer_distance = Inches(0.3)
    styles = doc.styles
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "List Bullet"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        if style.element.rPr is not None:
            color = style.element.rPr.find(qn("w:color"))
            if color is not None:
                for attr in ("themeColor", "themeTint", "themeShade"):
                    color.attrib.pop(qn(f"w:{attr}"), None)
    normal = styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name, size in (("Title", 28), ("Subtitle", 16), ("Heading 1", 20), ("Heading 2", 15), ("Heading 3", 12)):
        styles[name].font.size = Pt(size)
        styles[name].paragraph_format.keep_with_next = True
        styles[name].paragraph_format.space_before = Pt(12 if name.startswith("Heading") else 0)
        styles[name].paragraph_format.space_after = Pt(6)
    styles["List Bullet"].paragraph_format.space_after = Pt(3)
    styles["List Bullet"].paragraph_format.line_spacing = 1.05
    doc.core_properties.title = DATA["title"]
    doc.core_properties.subject = "Research foundation and complete editorial blueprint"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "AI engineering, software engineering, textbook, editorial blueprint"
    doc.add_paragraph(DATA["title"], "Title")
    doc.add_paragraph(DATA["subtitle"], "Subtitle")
    p = doc.add_paragraph("Research foundation and editorial blueprint")
    p.runs[0].bold = True
    doc.add_paragraph("September 5, 2026")
    doc.add_paragraph("A practical textbook for experienced software engineers who are new to machine learning. The book develops the ability to design, measure, constrain, and operate systems built around foundation models.")
    doc.add_paragraph("The blueprint specifies 30 chapters in eight parts, with an estimated 500 pages of main text. A fictional support application provides the recurring project, supported by focused side labs and an evidence-based capstone.")
    doc.add_heading("Review guide", 2)
    for text in ("Research foundation explains the evidence and boundaries behind the curriculum.",
                 "Detailed chapter outline provides topics, prerequisites, workshops, failure exercises, and completion evidence.",
                 "Project and assessment design defines the recurring application, milestones, educational targets, and companion material.",
                 "Source notes preserve attribution, publication context, and limits of the cited evidence."):
        doc.add_paragraph(text, "List Bullet")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Page ").font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for text in ((ROOT / "research/report-source.md").read_text(), outline,
                 (ROOT / "project-and-assessment.md").read_text(), source_notes()):
        add_markdown(doc, text)
    # Remove inherited decorative title rules as well as any direct paragraph rules.
    for element in list(doc.styles.element.iter(qn("w:pBdr"))) + list(doc.element.iter(qn("w:pBdr"))):
        element.getparent().remove(element)
    output = ROOT / "becoming-an-ai-engineer-blueprint.docx"
    doc.save(output)
    # Check final OOXML for expected content and accessible link relationships.
    from zipfile import ZipFile
    from lxml import etree
    with ZipFile(output) as archive:
        xml = etree.fromstring(archive.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        visible = " ".join(xml.xpath("//w:t/text()", namespaces=ns))
        assert all(c["title"] in visible for c in DATA["chapters"])
        assert "turn15view" not in visible and "<proposed_plan>" not in visible
        assert len(xml.xpath("//w:hyperlink", namespaces=ns)) >= len(SOURCES)
    return output


if __name__ == "__main__":
    validate()
    outline = outline_text()
    (ROOT / "book-outline.md").write_text(outline)
    output = create_doc(outline)
    print(json.dumps({"chapters": 30, "parts": 8, "planned_main_text_pages": 500,
                      "sources": len(SOURCES), "outline_words": len(outline.split()),
                      "docx": str(output), "bytes": output.stat().st_size}, indent=2))
